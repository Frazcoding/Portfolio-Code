"""Evaluate every formula in the capex/opex workbook with the `formulas` engine (no Excel/LibreOffice here), report
errors and print the key results per scenario for comparison with the Python model.

    python bplan_workbook_check.py
"""
import os
import sys
import warnings
from pathlib import Path

import numpy as np
import openpyxl

warnings.filterwarnings("ignore")
import formulas  # noqa: E402

PATH = Path(os.environ.get("CEA_XLSX_OUT") or Path(__file__).resolve().parent.parent / "business_plan" / "Glasshouse_Capex_Opex_Model.xlsx")


class WorkbookCheck:
    LABELS = {"Capex": ["Structure total", "Heat pump", "Heat system total", "GREENHOUSE", "ENERGY-ROUTE", "TOTAL DAY-ONE",
                        "Greenhouse capex per m²"],
              "Energy": ["LCOE", "Blended", "Turbine business"],
              "Opex": ["Revenue", "Electricity", "TOTAL RUNNING", "Running cost per kg"],
              "Results": ["Reduction", "LEVELISED", "of which capital", "Margin", "Break-even", "Greenhouse NPV"],
              "Scale": ["Greenhouse capex per m²", "Change vs 1 ha"],
              "Prices": ["Total"], "Inputs": ["Cucumber price"]}

    def __init__(self, path=PATH):
        self.path = path
        model = formulas.ExcelModel().loads(str(path)).finish()
        self.sol = model.calculate()
        self.vals, self.errors = {}, []
        for k, v in self.sol.items():
            try:
                a = np.asarray(v.value).ravel()
            except Exception:
                continue
            for item in a:
                if str(item).startswith("#"):
                    self.errors.append((k, str(item)))
            self.vals[k.upper()] = a[0] if a.size else None

    def value(self, sheet, col, row):
        return self.vals.get(f"'[{self.path.name.upper()}]{sheet.upper()}'!{col}{row}")

    def report(self):
        print(f"formula cells evaluated: {len(self.vals)}, errors: {len(self.errors)}")
        for e in self.errors[:30]:
            print("  ERROR", e)
        wb = openpyxl.load_workbook(self.path)
        for sheet, labels in self.LABELS.items():
            ws = wb[sheet]
            for r in range(1, ws.max_row + 1):
                lab = ws.cell(r, 1).value
                if lab and any(lb in str(lab) for lb in labels):
                    cols = "BCDEFG" if sheet in ("Scale", "Inputs", "Prices") else "CDEFG"
                    row = []
                    for c in cols:
                        v = self.value(sheet, c, r)
                        row.append(round(float(v), 3) if isinstance(v, (int, float, np.floating)) else v)
                    print(f"{sheet[:8]:9}{str(lab)[:50]:51}{row}")
        return len(self.errors)


if __name__ == "__main__":
    sys.exit(1 if WorkbookCheck().report() else 0)
